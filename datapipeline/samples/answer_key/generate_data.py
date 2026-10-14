"""
generate_data.py  (ANSWER KEY - do not share with the agent)

Generates ONE synthetic set of Medicare Part A style claims and writes it in two
structures:
  - agent_inputs/legacy/legacy_data.xlsx      (sheets PROV, CLM, STS)
  - agent_inputs/new_system/new_data.xlsx     (sheets provider, claim, claim_status)

All data is fictitious. NPIs and beneficiary IDs are randomly generated and do not
belong to real providers or people. No PHI.

Usage:  python generate_data.py
"""
import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

SEED = 42
N_CLAIMS = 120
ROOT = Path(__file__).resolve().parent.parent
LEGACY_XLSX = ROOT / "agent_inputs" / "legacy" / "legacy_data.xlsx"
NEW_XLSX = ROOT / "agent_inputs" / "new_system" / "new_data.xlsx"

rng = random.Random(SEED)

# ---------------------------------------------------------------- reference data
PROVIDERS = [
    ("1932456781", "Hudson Valley Regional Hospital", "NY"),
    ("1845023967", "Lakeshore Community Medical Center", "IL"),
    ("1760394825", "Granite State Surgical Hospital", "NH"),
    ("1598203746", "North Star Outpatient Clinic", "MN"),
    ("1487362019", "Bay Colony Health System", "MA"),
    ("1376029485", "Evergreen Hospice and Palliative Care", "WI"),
    ("1265803947", "Riverside Skilled Nursing Facility", "CT"),
    ("1154927306", "Pine Ridge Rural Health Center", "ME"),
    ("1043685291", "Metro Heart and Vascular Institute", "NY"),
    ("1932087465", "Prairie View Medical Center", "IL"),
    ("1821746093", "Coastal Rehabilitation Hospital", "RI"),
    ("1710359824", "Twin Lakes Community Hospital", "MN"),
]

STATUSES = [  # code, description, is_payable
    ("P", "PAID", "Y"),
    ("A", "ADJUSTED", "Y"),
    ("R", "REJECTED", "N"),
    ("V", "VOIDED", "N"),
    ("D", "DENIED", "N"),
]

# bill type -> (weight, charge range, span days range)
BILL_TYPES = {
    "111": (35, (8000, 45000), (2, 12)),   # hospital inpatient, admit-through-discharge
    "131": (38, (300, 4500), (0, 0)),      # hospital outpatient
    "137": (5, (300, 4500), (0, 0)),       # hospital outpatient, replacement claim
    "811": (11, (3000, 9000), (10, 29)),   # hospice (non-hospital based)
    "121": (11, (1500, 6000), (0, 3)),     # hospital inpatient Part B
}

MBI_LETTERS = "ACDEFGHJKMNPQRTUVWXY"   # MBI excludes S, L, O, I, B, Z


def fake_mbi():
    d = lambda: rng.choice("123456789")
    l = lambda: rng.choice(MBI_LETTERS)
    a = lambda: rng.choice(MBI_LETTERS + "0123456789")
    return d() + l() + a() + d() + l() + a() + d() + l() + l() + d() + d()


BENES = [fake_mbi() for _ in range(70)]


def pick_weighted(options):
    keys = list(options)
    return rng.choices(keys, weights=[options[k][0] for k in keys])[0]


def money(x):
    return round(x + 1e-9, 2)


# ---------------------------------------------------------------- claims
claims = []
for i in range(1, N_CLAIMS + 1):
    npi = rng.choice(PROVIDERS)[0]
    bill_typ = pick_weighted(BILL_TYPES)
    _, (cmin, cmax), (smin, smax) = BILL_TYPES[bill_typ]
    svc_from = date(2025, 1, 1) + timedelta(days=rng.randint(0, 180))
    svc_thru = svc_from + timedelta(days=rng.randint(smin, smax))
    received = svc_thru + timedelta(days=rng.randint(3, 60))
    status = rng.choices(["P", "A", "R", "V", "D"], weights=[74, 6, 7, 5, 8])[0]
    charges = money(rng.uniform(cmin, cmax))
    claims.append(dict(claim_id=f"C25{i:06d}", npi=npi, bene=rng.choice(BENES),
                       bill_typ=bill_typ, svc_from=svc_from, svc_thru=svc_thru,
                       received=received, status=status, charges=charges))

# ---- engineered edge cases (fixed positions so the answer key is stable) ----
def force(idx, **kw):
    claims[idx].update(kw)

# Late filing by both legacy and new rules (> 365 days after through date)
force(4, status="P", received=claims[4]["svc_thru"] + timedelta(days=402))
force(17, status="P", received=claims[17]["svc_thru"] + timedelta(days=380))
force(58, status="A", received=claims[58]["svc_thru"] + timedelta(days=455))

# Rule-change boundary: late by FROM date (legacy drops) but on time by THROUGH date (new keeps)
for idx, span, lag_from in [(9, 10, 369), (33, 8, 368), (71, 20, 375)]:
    frm = claims[idx]["svc_from"]
    force(idx, bill_typ="111" if span < 15 else "811", status="P",
          svc_thru=frm + timedelta(days=span), received=frm + timedelta(days=lag_from))
    claims[idx]["charges"] = money(rng.uniform(9000, 30000))

# Exact boundaries on the through date: 365 = on time, 366 = late
force(44, bill_typ="131", status="P", svc_thru=claims[44]["svc_from"],
      received=claims[44]["svc_from"] + timedelta(days=365))
force(88, bill_typ="131", status="P", svc_thru=claims[88]["svc_from"],
      received=claims[88]["svc_from"] + timedelta(days=366))

# Missing received date (legacy silently keeps these; new sends to exceptions)
force(26, status="P", received=None)
force(95, status="P", received=None)

# Denied claims are guaranteed present (legacy keeps them; new excludes them)
force(12, status="D")
force(63, status="D")
# A claim that is both not payable AND late -> reason priority test
force(101, status="R", received=claims[101]["svc_thru"] + timedelta(days=420))

# ---- amounts (after statuses are final) ----
for c in claims:
    if c["status"] in ("R", "V", "D"):
        c.update(allowed=0.0, deductible=0.0, coinsurance=0.0)
        continue
    allowed = money(c["charges"] * rng.uniform(0.35, 0.75))
    bt = c["bill_typ"]
    if bt == "111":
        ded = 1676.00 if rng.random() < 0.4 else 0.0     # Part A deductible (2025)
        coins = 0.0
    elif bt == "811":
        ded, coins = 0.0, 0.0
    else:  # outpatient / Part B style
        ded = money(min(allowed, rng.choice([0.0, 0.0, 257.00, rng.uniform(40, 257)])))
        coins = money((allowed - ded) * 0.20)
    c.update(allowed=allowed, deductible=ded, coinsurance=coins)

# ---------------------------------------------------------------- legacy structure
def yyyymmdd(d):
    return d.strftime("%Y%m%d") if d else None

prov_legacy = pd.DataFrame(PROVIDERS, columns=["PROV_NPI", "PROV_NM", "PROV_ST"])
prov_legacy["PROV_NM"] = prov_legacy["PROV_NM"].str.upper()
sts_legacy = pd.DataFrame([(s[0], s[1]) for s in STATUSES], columns=["STS_CD", "STS_DESC"])
clm_legacy = pd.DataFrame([{
    "CLM_ID": c["claim_id"], "PROV_NPI": c["npi"], "BENE_ID": c["bene"],
    "BILL_TYP": c["bill_typ"], "SVC_FROM_DT": yyyymmdd(c["svc_from"]),
    "SVC_THRU_DT": yyyymmdd(c["svc_thru"]), "RCVD_DT": yyyymmdd(c["received"]),
    "CLM_CHRG_AMT": c["charges"], "CLM_ALWD_AMT": c["allowed"],
    "PT_RESP_AMT": money(c["deductible"] + c["coinsurance"]),
    "CLM_STS_CD": c["status"],
} for c in claims])

# ---------------------------------------------------------------- new structure
provider_new = pd.DataFrame(PROVIDERS, columns=["npi", "provider_name", "state_code"])
status_new = pd.DataFrame(STATUSES, columns=["status_code", "status_description", "is_payable"])
status_new["status_description"] = status_new["status_description"].str.title()
claim_new = pd.DataFrame([{
    "claim_id": c["claim_id"], "npi": c["npi"], "beneficiary_id": c["bene"],
    "facility_type_code": c["bill_typ"][0], "bill_classification_code": c["bill_typ"][1],
    "claim_frequency_code": c["bill_typ"][2],
    "service_from_date": c["svc_from"], "service_through_date": c["svc_thru"],
    "received_date": c["received"],
    "billed_amount": c["charges"], "allowed_amount": c["allowed"],
    "deductible_amount": c["deductible"], "coinsurance_amount": c["coinsurance"],
    "status_code": c["status"],
} for c in claims])


# ---------------------------------------------------------------- write Excel
def write_xlsx(path, sheets, date_cols=()):
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    hdr_fill = PatternFill("solid", start_color="D9E1F2")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
                if cell.row == 1:
                    cell.fill = hdr_fill
                elif ws.cell(1, cell.column).value in date_cols and cell.value is not None:
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, str):
                    cell.number_format = "@"   # keep codes/IDs/text dates as text
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 42)
        ws.freeze_panes = "A2"
    wb.save(path)


write_xlsx(LEGACY_XLSX, {"PROV": prov_legacy, "CLM": clm_legacy, "STS": sts_legacy})
write_xlsx(NEW_XLSX, {"provider": provider_new, "claim": claim_new, "claim_status": status_new},
           date_cols=("service_from_date", "service_through_date", "received_date"))

print(f"Wrote {LEGACY_XLSX.name}: {len(clm_legacy)} claims, {len(prov_legacy)} providers")
print(f"Wrote {NEW_XLSX.name}: {len(claim_new)} claims, {len(provider_new)} providers")
print("Status mix:", clm_legacy["CLM_STS_CD"].value_counts().to_dict())
print("Bill type mix:", clm_legacy["BILL_TYP"].value_counts().to_dict())
